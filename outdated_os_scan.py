#!/usr/bin/env python3
# =============================================================================
# outdated_os_scan.py
#
# AUTHORIZED INTERNAL SECURITY ASSESSMENT TOOL
# Scope: OS/kernel version fingerprinting and deprecated/EOL-OS discovery on
#        internal company networks. Same two-phase masscan + worker-pool
#        architecture as ssh_vuln_scan.py / telnet_spray.py / sslspray.py /
#        quantum_readiness_spray.py.
#
# -----------------------------------------------------------------------------
# WHAT THIS SCRIPT DOES
# -----------------------------------------------------------------------------
# Phase 1 (Discovery): masscan sweeps the configured subnets for a short list
#   of commonly-open ports (MASSCAN_PORTS) to find live hosts fast - same
#   rationale as every sibling script: the configured scope includes a /8,
#   and pointing per-host nmap -O at that much address space directly would
#   dominate the whole run.
# Phase 2 (OS fingerprint): for every host discovered in Phase 1, a pool of
#   worker threads each run one
#       nmap -O -sV -Pn -n --max-rate <rate> --host-timeout <timeout>
#            -oX <tmpfile> <ip>
#   subprocess against that single host (one nmap process per host, not
#   batched - see WHY THIS REWRITE EXISTS for why that changed). nmap's XML
#   <os><osmatch> candidates are parsed, filtered by MIN_ACCURACY, and
#   checked against the DEPRECATED_OS_RULES table.
#
# -----------------------------------------------------------------------------
# WHY THIS REWRITE EXISTS
# -----------------------------------------------------------------------------
# This was ported off the original `outdatedOS.py` (still at the repo root,
# untouched by this rewrite - a separate process handles retiring it), which
# had two problems this script fixes:
#
# 1. UNDERCOUNTING BUG ("grab all the outdated OS's so we get an accurate
#    count"): the original's parse_outdated() did
#        osmatch = os_elem.find("osmatch")
#    Python's Element.find() returns only the FIRST <osmatch> child. nmap's
#    XML lists multiple <osmatch> candidates per host (nmap can report up to
#    ~10), already sorted by accuracy descending. A host whose single best
#    guess ISN'T a string in DEPRECATED_OS_KEYWORDS, but whose 2nd- or
#    3rd-best guess (still at/above MIN_ACCURACY) IS one, was entirely
#    invisible to the original tool - undercounting deprecated hosts.
#    Concrete example - nmap returns, for one host:
#        osmatch 1: "Linux 5.4"        accuracy=93   (not deprecated)
#        osmatch 2: "Linux 2.6.32"     accuracy=91   (CRITICAL - EOL kernel)
#        osmatch 3: "FreeBSD 10.1"     accuracy=90   (not deprecated)
#    osmatch[0] alone never matches anything in the deprecated-OS table, so
#    the original tool silently dropped this host even though a second,
#    still-high-confidence guess flags it CRITICAL. This script instead
#    walks every <osmatch> child in document order (already accuracy-sorted
#    by nmap) and takes the FIRST one that is BOTH at/above MIN_ACCURACY AND
#    matches a DEPRECATED_OS_RULES keyword - so the example above now
#    correctly reports "Linux 2.6.32" as CRITICAL. Every osmatch examined
#    (deprecated-matching or not) is also captured verbatim in the
#    "Alternate OS Guesses" column as evidence, so a human reviewer can see
#    what else nmap considered.
# 2. CONSOLE-ONLY PROGRESS: the original used the `rich` library for a
#    multi-row live worker display. Every sibling script in this directory
#    instead uses a single dependency-free progress line (draw_progress_line
#    / finish_progress_line / render_bar) plus a file+console logger
#    (setup_logging / ProgressAwareHandler) - ported verbatim here too, so
#    this script has the same zero-extra-dependency footprint (besides the
#    optional openpyxl) as its siblings.
#
# Also changed, deliberately, from a batched to a per-host Phase 2: the
# original ran one nmap invocation per 64-host batch (FINGERPRINT_BATCH).
# Every sibling script's Phase 2 instead runs one subprocess per host (see
# ssh_vuln_scan.py's _run_nmap_single_host / scan_one_host), which gives
# per-host progress feedback and per-host retry granularity instead of
# having to re-run (or silently lose) an entire 64-host batch on one bad
# invocation.
#
# -----------------------------------------------------------------------------
# CATEGORY DEFINITIONS (severity buckets - every report-body row belongs to
# exactly one of these three; see LIMITATIONS for why there's no "clean"
# bucket here, unlike the 5-bucket sibling scripts)
# -----------------------------------------------------------------------------
# [red]    Critical - Fully Unsupported (Legacy OS) - no vendor patches
#          available at all, for anything, including already-public CVEs
#          (e.g. Windows XP, Windows Server 2008, Linux 2.6.x, Solaris,
#          HP-UX).
# [orange] High - End of Life - vendor support has ended or is ending
#          imminently (e.g. Windows 7, Windows Server 2012, Linux 3.x
#          kernels, VMware ESXi 6.x).
# [yellow] Watch - Recently / Soon End of Life - still inside its support
#          window at time of writing, but widely deployed and close enough
#          to EOL to flag now (e.g. Windows 10, EOL 2025-10-14).
#
# Separately, the Overview sheet's Scan Summary tracks (NOT as report-body
# categories, since this report only ever lists deprecated hosts, matching
# the original tool's scope): total live hosts discovered in Phase 1, hosts
# that returned ANY osmatch data ("fingerprinted"), and hosts that returned
# none at all ("not fingerprinted" / inconclusive - nmap -O failed to get
# enough signal).
#
# -----------------------------------------------------------------------------
# NON-DESTRUCTIVE / SAFETY GUARANTEES
# -----------------------------------------------------------------------------
# nmap -O/-sV fingerprinting only sends standard TCP/IP-stack and service
# probe packets to infer OS/service identity from response behavior - no
# authentication is attempted, no exploitation, no data modification. Same
# as every sibling script, masscan's own Phase 1 probes are plain SYN
# packets.
#
# THIS TOOL MUST ONLY BE RUN AGAINST NETWORKS YOU ARE EXPLICITLY AUTHORIZED
# TO ASSESS. Confirm written authorization / an active engagement scope
# before running this script.
#
# -----------------------------------------------------------------------------
# LIMITATIONS AND ASSUMPTIONS
# -----------------------------------------------------------------------------
#   - Requires `masscan` on PATH (unless --skip-masscan with a valid
#     --masscan-output-file) AND `nmap` on PATH (always - no nmap -sn
#     fallback for Phase 1, unlike the original script, and no Phase-2
#     alternative to nmap -O/-sV). This is a deliberate scope decision to
#     match the sibling scripts' own convention: none of them carry an
#     nmap -sn fallback, and adding one back here just for this tool would
#     make it the odd one out.
#   - `openpyxl` is only needed for the .xlsx step; its absence degrades to
#     CSV-only rather than failing the run (imported lazily, never at
#     module level).
#   - A masscan hit on a probed port only proves that port is open / the
#     host is alive - NOT that nmap -O will succeed at fingerprinting it.
#     Firewalling of the OTHER probe traffic nmap sends (closed-port resets,
#     ICMP, etc.) can still leave a host "not fingerprinted" even though
#     Phase 1 found it alive.
#   - nmap -O itself is a heuristic best-effort TCP/IP-stack fingerprint,
#     even for a match reported above MIN_ACCURACY - treat every finding as
#     a strong lead to verify, not an absolute fact, same as any other
#     passive fingerprinting technique.
#   - EOL dates/notes in DEPRECATED_OS_RULES are approximate, public-
#     knowledge context captured at the time this script was written - they
#     are NOT guaranteed current. Re-verify against the vendor's own current
#     lifecycle page before quoting a specific date in a client-facing
#     deliverable.
#   - Windows Server 2016+ is deliberately NOT flagged - its extended
#     support had not yet expired at the time this script (and the original
#     it was ported from) was written. Re-evaluate as that changes.
#   - Reverse DNS depends on corporate DNS infrastructure; failures resolve
#     to "" (blank) and do not stop the scan.
#   - Default --rate is 25000 pps, matching the sibling scripts - the
#     original outdatedOS.py this was ported from used a deliberately
#     conservative, network-team-coordinated 2000 pps instead. Confirm the
#     configured rate is appropriate for your specific engagement/network
#     before a live run (same caution the original script's own comment
#     carried); Phase 1 here is the same masscan SYN sweep the sibling
#     tools already run against this SUBNETS list, not a different load.
# =============================================================================

import argparse
import csv
import ipaddress
import logging
import os
import re
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from concurrent.futures import TimeoutError as FutureTimeoutError
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, List, Optional, Tuple
from xml.etree import ElementTree as ET

# =============================================================================
# CONFIGURATION
# =============================================================================

SCRIPT_DIR: str = os.path.dirname(os.path.abspath(__file__))

# Same subnets in scope as the sibling sweep tools. Edit this list to change
# scope.
SUBNETS: List[str] = [
    "1.0.0.0/16",
    "2.0.0.0/16",
    "3.0.0.0/16",
    "4.0.0.0/16",
    "5.0.0.0/16",
    # "6.0.0.0/16",
    "7.0.0.0/16",
    "8.0.0.0/12",
    "9.0.0.0/8",
]

# Ports masscan probes in Phase 1 (ported verbatim from the original
# outdatedOS.py's MASSCAN_PORTS - a short list of commonly-open ports that
# gives nmap -O something to anchor an open/closed-port TCP/IP fingerprint
# on, same rationale the original script documented).
MASSCAN_PORTS = "21,22,23,25,53,80,110,135,139,143,443,445,3389,5900,8080"

DEFAULT_WORKERS = 10    # lower than ssh_vuln_scan.py's 16: -O -sV is heavier
                         # per-host than NSE-script-only work (OS fingerprint
                         # probing plus a service-version probe on every open
                         # port it finds), so each worker costs more wall-
                         # clock time per host.
DEFAULT_RATE = 25000     # masscan packets/sec for Phase 1, same default as
                         # the sibling tools. Also passed to nmap's own
                         # --max-rate in Phase 2 (the CLI only exposes one
                         # --rate flag, matching the rest of the flag set
                         # spec'd for this tool) - this is a ceiling, not a
                         # target, so reusing the same (large) number is
                         # harmless: a single nmap -O/-sV run against one
                         # host never generates anywhere near that much
                         # traffic on its own.
                         #
                         # NOTE: the original outdatedOS.py this was ported
                         # from fixed Phase 1 at a deliberately conservative
                         # 2000 pps with an explicit "coordinate with your
                         # network team" comment. This default was raised to
                         # match the sibling scripts' already-established,
                         # already-in-use-against-this-same-SUBNETS-list rate
                         # - Phase 1 here is the identical masscan SYN sweep
                         # those tools already run, not a materially
                         # different load. Still: confirm the configured
                         # --rate is appropriate for your specific engagement
                         # / network before a live run, same as any scanner.
DEFAULT_HOST_TIMEOUT = "60s"  # longer than ssh_vuln_scan's 30s default: -O's
                               # TCP/IP stack fingerprint sends more probe
                               # packets and waits on more timing windows
                               # than an NSE script does, so a hung/filtered
                               # host needs more rope before being written
                               # off, while still bounding a worker slot.
DEFAULT_RETRIES = 1
MIN_ACCURACY = 90        # module-level default; overridable via --min-accuracy.
                         # Ignore osmatch guesses below this confidence (%).

CSV_FIELDS = ["Scan Date", "IP Address", "Hostname", "Subnet Range", "OS Name",
              "OS Family", "Accuracy", "Risk Category", "Alternate OS Guesses",
              "Notes", "Meaning", "Scan Status"]

# -----------------------------------------------------------------------------
# OS risk buckets (mirrors telnet_spray.py's AUTH_* dict pattern, renamed
# OS_RISK_*). Only 3 buckets, not 5 - see module docstring CATEGORY
# DEFINITIONS for why there is no "clean" bucket here.
# -----------------------------------------------------------------------------
OS_RISK_CRITICAL = "Critical - Fully Unsupported (Legacy OS)"
OS_RISK_HIGH = "High - End of Life"
OS_RISK_WATCH = "Watch - Recently / Soon End of Life"

OS_RISK_ORDER = [OS_RISK_CRITICAL, OS_RISK_HIGH, OS_RISK_WATCH]


def normalize_risk_category(cat: str) -> str:
    """The one place 'what do we do with an unrecognized Risk Category'
    policy lives. Previously reimplemented independently at three call
    sites (CSV reload, compute_category_stats, build_workbook's row-render
    loop) with three different defaulting styles - changing the policy (or
    widening it to a 4th bucket) meant finding and editing all three."""
    return cat if cat in OS_RISK_ORDER else OS_RISK_HIGH


OS_RISK_EMOJI = {
    OS_RISK_CRITICAL: "\U0001F534",  # red circle
    OS_RISK_HIGH: "\U0001F7E0",      # orange circle
    OS_RISK_WATCH: "\U0001F7E1",     # yellow circle
}

OS_RISK_FILL_HEX = {
    OS_RISK_CRITICAL: "FFC7CE",
    OS_RISK_HIGH: "FCE4D6",
    OS_RISK_WATCH: "FFEB9C",
}

OS_RISK_FONT_HEX = {
    OS_RISK_CRITICAL: "922B21",
    OS_RISK_HIGH: "784212",
    OS_RISK_WATCH: "7D6608",
}

# One-line, per-row "Meaning" text.
OS_RISK_MEANING = {
    OS_RISK_CRITICAL: ("CRITICAL - This OS/kernel is fully out of vendor support, with no patches "
                        "available for any vulnerability, including already-public CVEs. Prioritize "
                        "migration, isolation, or compensating network controls immediately."),
    OS_RISK_HIGH: ("HIGH - End of life. Vendor patches have stopped (or are stopping imminently) for "
                   "this OS version. Treat as unsupported for any new or critical CVE and prioritize "
                   "an upgrade path."),
    OS_RISK_WATCH: ("WATCH - Still inside its support window at time of writing, but a widely-deployed, "
                     "soon-to-be (or recently) end-of-life OS worth tracking now rather than waiting "
                     "for the EOL date to arrive unplanned."),
}

# Longer, paragraph-length text for the Overview sheet's category table.
OS_RISK_DESCRIPTION = {
    OS_RISK_CRITICAL: (
        "No vendor patches are available at all for this OS/kernel, for anything, including "
        "vulnerabilities that are already public knowledge. Examples: Windows 95/98/2000/XP/Vista, "
        "Windows Server 2000/2003/2008, Linux kernels 2.0-2.6.x, Solaris/SunOS, old FreeBSD/AIX, "
        "HP-UX, VMware ESXi 5.x. Highest-priority remediation target."
    ),
    OS_RISK_HIGH: (
        "Vendor support has ended, or is ending imminently. Examples: Windows 7, Windows Server 2012, "
        "Linux 3.x kernels (including RHEL 7's 3.10), VMware ESXi 6.x, generic 'Mac OS X 10.x' "
        "fingerprints. Treat as unsupported for new/critical CVEs; plan an upgrade."
    ),
    OS_RISK_WATCH: (
        "Still inside its support window at the time this scan ran, but widely deployed and close "
        "enough to its own end-of-life date to flag now rather than after the fact. Currently just "
        "Windows 10 (EOL 2025-10-14)."
    ),
}

# -----------------------------------------------------------------------------
# Deprecated-OS rules table: (keyword_substring, family, risk_category,
# eol_note). Evaluated in order - first substring match (case-insensitive)
# wins. EOL dates/notes are approximate public-knowledge context at the time
# this was written - see LIMITATIONS in the module docstring.
# -----------------------------------------------------------------------------
DEPRECATED_OS_RULES: List[Tuple[str, str, str, str]] = [
    # --- Windows client ---
    ("Windows 95", "Windows client", OS_RISK_CRITICAL, "EOL 2001-12-31"),
    ("Windows 98", "Windows client", OS_RISK_CRITICAL, "EOL 2006-07-11"),
    ("Windows Millennium", "Windows client", OS_RISK_CRITICAL,
     "EOL 2006-07-11 (nmap also reports \"Windows Me\")"),
    ("Windows 2000", "Windows client", OS_RISK_CRITICAL, "EOL 2010-07-13"),
    ("Windows XP", "Windows client", OS_RISK_CRITICAL, "EOL 2014-04-08"),
    ("Windows Vista", "Windows client", OS_RISK_CRITICAL, "EOL 2017-04-11"),
    ("Windows 7", "Windows client", OS_RISK_HIGH, "EOL 2020-01-14"),
    ("Windows 8", "Windows client", OS_RISK_HIGH,
     "EOL 2016-01-12 (8.0) / 2023-01-10 (8.1) - substring also catches \"Windows 8.1\""),
    ("Windows 10", "Windows client", OS_RISK_WATCH, "EOL 2025-10-14"),
    # --- Windows Server ---
    ("Windows Server 2000", "Windows Server", OS_RISK_CRITICAL, "EOL 2010-07-13"),
    ("Windows Server 2003", "Windows Server", OS_RISK_CRITICAL, "EOL 2015-07-14"),
    ("Windows Server 2008", "Windows Server", OS_RISK_CRITICAL,
     "EOL 2020-01-14 (substring also catches \"2008 R2\")"),
    ("Windows Server 2012", "Windows Server", OS_RISK_HIGH,
     "EOL 2023-10-10 (substring also catches \"2012 R2\")"),
    # (Server 2016+ extended support had not expired at time of writing -
    # deliberately not flagged, see LIMITATIONS.)
    # --- Linux kernels (nmap fingerprints kernel version, NOT distro) ---
    ("Linux 2.0", "Linux", OS_RISK_CRITICAL, "kernel EOL, long unsupported"),
    ("Linux 2.2", "Linux", OS_RISK_CRITICAL, "kernel EOL"),
    ("Linux 2.4", "Linux", OS_RISK_CRITICAL, "kernel EOL"),
    ("Linux 2.6", "Linux", OS_RISK_CRITICAL, "kernel EOL (all 2.6.x)"),
    ("Linux 3.", "Linux", OS_RISK_HIGH,
     "kernel EOL (3.0-3.19, incl. RHEL 7's 3.10 - trailing dot avoids matching \"Linux 30\"-style false positives)"),
    # --- Other Unix / hypervisors ---
    ("Solaris", "Solaris / SunOS", OS_RISK_CRITICAL, "vendor EOL"),
    ("SunOS", "Solaris / SunOS", OS_RISK_CRITICAL, "vendor EOL"),
    ("FreeBSD 4", "FreeBSD", OS_RISK_CRITICAL, "EOL"),
    ("FreeBSD 5", "FreeBSD", OS_RISK_CRITICAL, "EOL"),
    ("FreeBSD 6", "FreeBSD", OS_RISK_CRITICAL, "EOL"),
    ("FreeBSD 7", "FreeBSD", OS_RISK_CRITICAL, "EOL"),
    ("AIX 5", "AIX", OS_RISK_CRITICAL, "EOL"),
    ("AIX 6", "AIX", OS_RISK_CRITICAL, "EOL"),
    ("HP-UX", "HP-UX", OS_RISK_CRITICAL,
     "EOL (generic match - specific version unknown from nmap fingerprint)"),
    ("VMware ESXi 5", "VMware ESXi", OS_RISK_CRITICAL, "EOL 2020-03-01"),
    ("VMware ESXi 6", "VMware ESXi", OS_RISK_HIGH, "EOL (6.0/6.5/6.7 all EOL)"),
    ("Mac OS X", "Mac OS X", OS_RISK_HIGH,
     "generic \"Mac OS X 10.x\" match - modern macOS reports as \"macOS\", not \"Mac OS X\""),
]


def classify_deprecated(os_name: str) -> Optional[Tuple[str, str, str]]:
    """Return (family, risk_category, eol_note) for the first
    DEPRECATED_OS_RULES entry whose keyword is a substring of os_name
    (case-insensitive), or None if nothing matches."""
    lower = (os_name or "").lower()
    for keyword, family, risk_category, eol_note in DEPRECATED_OS_RULES:
        if keyword.lower() in lower:
            return family, risk_category, eol_note
    return None


# =============================================================================
# TEXT SANITIZATION (ported verbatim from telnet_spray.py)
# =============================================================================
# OS Name / Alternate OS Guesses come straight out of nmap's <osmatch name=...>
# attribute (nmap's own fingerprint-database string, not raw attacker bytes,
# but still untrusted external input as far as this script is concerned), and
# Hostname comes from reverse DNS - same category of "don't trust it to be
# well-formed or short" as the sibling scripts' banner/PTR fields.

_CONTROL_CHAR_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _sanitize_text(value: str, max_len: int = 2000) -> str:
    if not value:
        return ""
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = _CONTROL_CHAR_RE.sub("", value)
    value = value.strip()
    if len(value) > max_len:
        value = value[:max_len] + " ...(truncated)"
    return value


# CSV/Excel formula injection (CWE-1236): OS Name / Alternate OS Guesses /
# Hostname are all sourced from (or built from) external, untrusted input
# (nmap's fingerprint string, reverse DNS). Prefixing a leading quote onto
# any value starting with a formula-trigger character forces plain-text
# interpretation, matching the fix already applied in sslspray.py /
# telnet_spray.py.
_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@", "\t", "\r")


def _neutralize_formula(value):
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGER_CHARS):
        return "'" + value
    return value


# =============================================================================
# LOGGING / PROGRESS DISPLAY (ported near-verbatim from the sibling scripts)
# =============================================================================

_progress_lock = threading.Lock()
_last_progress_len = 0


class ProgressAwareHandler(logging.StreamHandler):
    def emit(self, record: logging.LogRecord) -> None:
        global _last_progress_len
        with _progress_lock:
            if _last_progress_len:
                sys.stdout.write("\r" + " " * _last_progress_len + "\r")
                sys.stdout.flush()
            super().emit(record)
            _last_progress_len = 0


def setup_logging(log_path: str) -> logging.Logger:
    logger = logging.getLogger("outdated_os_scan")
    logger.setLevel(logging.DEBUG)
    file_handler = logging.FileHandler(log_path, encoding="utf-8")
    file_handler.setLevel(logging.DEBUG)
    file_handler.setFormatter(logging.Formatter(
        "%(asctime)s [%(levelname)s] %(message)s", "%Y-%m-%d %H:%M:%S"))
    console_handler = ProgressAwareHandler(sys.stdout)
    console_handler.setLevel(logging.INFO)
    console_handler.setFormatter(logging.Formatter("%(message)s"))
    logger.addHandler(file_handler)
    logger.addHandler(console_handler)
    return logger


def draw_progress_line(line: str) -> None:
    global _last_progress_len
    with _progress_lock:
        pad = max(0, _last_progress_len - len(line))
        sys.stdout.write("\r" + line + (" " * pad))
        sys.stdout.flush()
        _last_progress_len = len(line)


def finish_progress_line() -> None:
    global _last_progress_len
    with _progress_lock:
        if _last_progress_len:
            sys.stdout.write("\n")
            sys.stdout.flush()
        _last_progress_len = 0


def fmt_elapsed(seconds: float) -> str:
    seconds = int(max(0, seconds))
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def render_bar(pct: Optional[float], width: int = 30) -> str:
    if pct is None:
        return "[" + "-" * width + "]  n/a"
    pct = max(0.0, min(100.0, pct))
    filled = int(width * pct / 100.0)
    return "[" + "#" * filled + "-" * (width - filled) + f"] {pct:5.1f}%"


# =============================================================================
# DEPENDENCY / VALIDATION HELPERS (mirrors the sibling scripts)
# =============================================================================

def check_external_tool(name: str) -> Optional[str]:
    return shutil.which(name)


def validate_subnets(raw_subnets: List[str], logger: logging.Logger) -> List[ipaddress.IPv4Network]:
    networks = []
    for entry in raw_subnets:
        try:
            networks.append(ipaddress.ip_network(entry, strict=False))
        except ValueError as exc:
            logger.error(f"Skipping invalid CIDR '{entry}': {exc}")
    return networks


def subnet_for_ip(ip: str, networks: List[ipaddress.IPv4Network]) -> str:
    try:
        addr = ipaddress.ip_address(ip)
    except ValueError:
        return "UNKNOWN"
    for net in networks:
        if addr in net:
            return str(net)
    return "UNKNOWN"


def _call_with_hard_timeout(func, timeout: float):
    """Runs func() with an ACTUAL enforced wall-clock timeout.
    socket.setdefaulttimeout() does NOT bound socket.gethostbyaddr() (or
    getfqdn()) -- confirmed by direct testing against an unresolvable
    address: it kept blocking for several seconds regardless of the
    requested timeout value (0.5/1.0/2.0s all measured ~5s). Those calls
    hit the OS's own blocking resolver directly and never consult Python's
    socket-level timeout at all - a well-known CPython gotcha, not a typo.
    A throwaway single-worker executor gives up and returns None if func()
    hasn't completed in time; shutdown(wait=False) means the orphaned
    thread is left to finish resolving (or for the OS resolver to time out
    on its own) in the background rather than blocking this call."""
    executor = ThreadPoolExecutor(max_workers=1)
    future = executor.submit(func)
    try:
        return future.result(timeout=timeout)
    except FutureTimeoutError:
        return None
    finally:
        executor.shutdown(wait=False)


def resolve_hostname(ip: str, timeout: float) -> str:
    def _lookup() -> str:
        try:
            name, _, _ = socket.gethostbyaddr(ip)
            return _sanitize_text(name)
        except (socket.herror, socket.gaierror, socket.timeout, OSError):
            return ""
    return _call_with_hard_timeout(_lookup, timeout) or ""


def parse_port_spec(spec: str) -> List[int]:
    ports = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        try:
            ports.append(int(part))
        except ValueError:
            continue
    return ports


# =============================================================================
# PHASE 1: MASSCAN DISCOVERY (ported near-verbatim from the sibling scripts)
# =============================================================================

class MasscanStatus:
    def __init__(self) -> None:
        self.lock = threading.Lock()
        self.percent: Optional[float] = None
        self.eta: str = ""

    def update_from_line(self, line: str) -> None:
        m = re.search(r"(\d+(?:\.\d+)?)%\s+done", line)
        eta_m = re.search(r"done,\s*([\d:]+)\s*remaining", line)
        with self.lock:
            if m:
                try:
                    self.percent = float(m.group(1))
                except ValueError:
                    pass
            if eta_m:
                self.eta = eta_m.group(1)


def _masscan_stderr_reader(proc: subprocess.Popen, status: MasscanStatus) -> None:
    buf = b""
    stream = proc.stderr
    if stream is None:
        return
    try:
        while True:
            chunk = stream.read(256)
            if not chunk:
                break
            buf += chunk
            while True:
                idx_r = buf.find(b"\r")
                idx_n = buf.find(b"\n")
                candidates = [i for i in (idx_r, idx_n) if i != -1]
                if not candidates:
                    break
                idx = min(candidates)
                line = buf[:idx].decode(errors="ignore").strip()
                buf = buf[idx + 1:]
                if line:
                    status.update_from_line(line)
    except (ValueError, OSError):
        pass


def build_masscan_command(masscan_path: str, subnets: List[str], rate: int,
                           output_file: str, interface: Optional[str],
                           ports: List[int]) -> List[str]:
    port_spec = "T:" + ",".join(str(p) for p in ports)
    cmd = [masscan_path, "-p", port_spec, "--rate", str(rate), "-oL", output_file]
    if interface:
        cmd += ["-e", interface]
    cmd += subnets
    return cmd


def parse_masscan_list_output(path: str, start_offset: int = 0) -> Tuple[List[Tuple[str, int, str]], int]:
    records: List[Tuple[str, int, str]] = []
    if not os.path.exists(path):
        return records, start_offset
    with open(path, "rb") as f:
        f.seek(start_offset)
        chunk = f.read()
    if not chunk:
        return records, start_offset
    last_newline = chunk.rfind(b"\n")
    if last_newline == -1:
        return records, start_offset
    usable, new_offset = chunk[:last_newline + 1], start_offset + last_newline + 1
    for raw_line in usable.split(b"\n"):
        line = raw_line.decode(errors="ignore").strip()
        if not line or line.startswith("#"):
            continue
        parts = line.split()
        if len(parts) < 5:
            continue
        status, proto, port_s, ip, _ts = parts[:5]
        if status != "open":
            continue
        try:
            port = int(port_s)
        except ValueError:
            continue
        records.append((ip, port, proto))
    return records, new_offset


def run_masscan_phase1(masscan_path: str, subnets: List[str], rate: int,
                        output_file: str, interface: Optional[str],
                        ports: List[int], logger: logging.Logger,
                        stop_event: threading.Event
                        ) -> List[Tuple[str, int, str]]:
    cmd = build_masscan_command(masscan_path, subnets, rate, output_file, interface, ports)
    logger.info("Phase 1 - DISCOVERY starting")
    logger.debug(f"Masscan command: {' '.join(cmd)}")

    try:
        proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    except FileNotFoundError:
        logger.error(f"masscan executable not found at '{masscan_path}'.")
        return []
    except PermissionError as exc:
        logger.error(f"Permission error launching masscan: {exc}. "
                      f"Masscan typically requires root/administrator privileges.")
        return []
    except OSError as exc:
        logger.error(f"Failed to launch masscan: {exc}")
        return []

    status = MasscanStatus()
    reader_thread = threading.Thread(target=_masscan_stderr_reader, args=(proc, status), daemon=True)
    reader_thread.start()

    start_time = time.time()
    live_count = 0
    offset = 0
    seen: set = set()

    def _drain_new_records() -> None:
        nonlocal offset, live_count
        new_records, offset = parse_masscan_list_output(output_file, offset)
        for ip, port, _proto in new_records:
            key = (ip, port)
            if key in seen:
                continue
            seen.add(key)
            live_count += 1

    try:
        while True:
            if stop_event.is_set():
                logger.warning("Interrupt received, terminating masscan...")
                proc.terminate()
                try:
                    proc.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    proc.kill()
                break

            retcode = proc.poll()
            _drain_new_records()

            elapsed = time.time() - start_time
            with status.lock:
                pct = status.percent
                eta = status.eta

            bar = render_bar(pct)
            eta_str = eta if eta else "n/a"
            line = (f"Phase 1 - DISCOVERY {bar} | Port hits: {live_count} | "
                     f"Elapsed: {fmt_elapsed(elapsed)} | ETA: {eta_str}")
            draw_progress_line(line)

            if retcode is not None:
                time.sleep(0.3)
                _drain_new_records()
                break
            time.sleep(0.5)
    finally:
        finish_progress_line()

    if proc.returncode not in (0, None) and not stop_event.is_set():
        logger.warning(f"masscan exited with return code {proc.returncode}. "
                        f"Results collected so far will still be used.")

    final_records, _ = parse_masscan_list_output(output_file, 0)
    logger.info(f"Phase 1 - DISCOVERY complete. Port hits: {live_count}, "
                f"elapsed: {fmt_elapsed(time.time() - start_time)}")
    return final_records


# =============================================================================
# PHASE 2: PER-HOST NMAP -O -sV WORKER POOL
# =============================================================================

@dataclass
class DiscoveredHost:
    ip: str
    hostname: str = ""
    subnet: str = "UNKNOWN"


@dataclass
class HostScanResult:
    ip: str
    hostname: str = ""
    subnet: str = "UNKNOWN"
    fingerprinted: bool = False
    deprecated: bool = False
    os_name: str = ""
    accuracy: str = ""
    family: str = ""
    risk_category: str = ""
    eol_note: str = ""
    alternate_guesses: str = ""
    notes: str = ""
    scan_status: str = "Completed"


@dataclass
class OutdatedOSRecord:
    scan_date: str
    ip: str
    hostname: str
    subnet: str
    os_name: str
    os_family: str
    accuracy: str
    risk_category: str
    alternate_guesses: str
    notes: str
    meaning: str
    scan_status: str = "Completed"


def record_to_row(r: OutdatedOSRecord) -> Dict[str, object]:
    return {
        "Scan Date": r.scan_date, "IP Address": r.ip, "Hostname": r.hostname,
        "Subnet Range": r.subnet, "OS Name": r.os_name, "OS Family": r.os_family,
        "Accuracy": r.accuracy, "Risk Category": r.risk_category,
        "Alternate OS Guesses": r.alternate_guesses, "Notes": r.notes,
        "Meaning": r.meaning, "Scan Status": r.scan_status,
    }


def _host_ip(host_elem) -> Optional[str]:
    for addr in host_elem.findall("address"):
        if addr.get("addrtype") in ("ipv4", "ipv6"):
            ip = addr.get("addr")
            if ip:
                return ip
    return None


def _host_hostname_from_xml(host_elem) -> str:
    hostnames_el = host_elem.find("hostnames")
    if hostnames_el is not None:
        chosen = hostnames_el.find("hostname[@type='PTR']")
        if chosen is None:
            chosen = hostnames_el.find("hostname")
        if chosen is not None:
            return chosen.get("name", "") or ""
    return ""


def evaluate_osmatches(osmatch_elems, min_accuracy: int
                        ) -> Tuple[bool, Optional[Tuple[str, str, str, str, str]], str]:
    """THE CORE FIX (see module docstring WHY THIS REWRITE EXISTS #1).

    Walks every <osmatch> child in document order (nmap already sorts them
    accuracy-descending) instead of only looking at the first one. Returns:
      fingerprinted - True iff nmap returned at least one osmatch at all.
      chosen        - (os_name, accuracy_str, family, risk_category,
                       eol_note) for the FIRST osmatch that is both
                       >= min_accuracy AND matches a DEPRECATED_OS_RULES
                       keyword, or None if no osmatch qualifies. Because
                       nmap's own list is already accuracy-sorted and we
                       stop updating `chosen` once set, this is always the
                       highest-accuracy qualifying match, not just whichever
                       is first in the list - we don't blindly take a lower-
                       accuracy match once a higher one already matched,
                       and we don't stop scanning just because osmatch[0]
                       wasn't a match.
      alt_str       - every osmatch examined (matching or not), as
                       "name@accuracy%" pairs joined with "; ", sanitized
                       and length-capped - evidence for a human reviewer.
    """
    fingerprinted = len(osmatch_elems) > 0
    chosen: Optional[Tuple[str, str, str, str, str]] = None
    alt_parts: List[str] = []
    for om in osmatch_elems:
        name = om.get("name", "") or ""
        acc_s = om.get("accuracy", "0") or "0"
        alt_parts.append(f"{name}@{acc_s}%")
        if chosen is not None:
            continue  # already found the highest-accuracy deprecated match;
                       # keep looping only to finish building alt_parts.
        try:
            acc = int(acc_s)
        except ValueError:
            continue
        if acc < min_accuracy:
            continue
        rule = classify_deprecated(name)
        if rule is not None:
            family, risk_category, eol_note = rule
            chosen = (name, acc_s, family, risk_category, eol_note)
    alt_str = _sanitize_text("; ".join(alt_parts), max_len=1500)
    return fingerprinted, chosen, alt_str


def parse_single_host_xml(xml_path: str, hostname_override: str, subnet_label: str,
                           min_accuracy: int) -> Optional[HostScanResult]:
    """Parses the single <host> element nmap wrote for one Phase-2 scan.
    Returns None if the XML is missing/unparseable or has no usable <host>
    (i.e. nmap didn't produce a result at all) - the caller treats that as
    retry-worthy. A host that parsed fine but has zero <osmatch> children is
    NOT None - it's a normal, deterministic "not fingerprinted" result."""
    try:
        tree = ET.parse(xml_path)
    except (ET.ParseError, FileNotFoundError, OSError):
        return None
    root = tree.getroot()
    host_elem = root.find("host")
    if host_elem is None:
        return None
    status = host_elem.find("status")
    if status is None or status.get("state") != "up":
        return None
    ip = _host_ip(host_elem)
    if ip is None:
        return None

    hostname = hostname_override or _sanitize_text(_host_hostname_from_xml(host_elem))

    os_elem = host_elem.find("os")
    osmatches = os_elem.findall("osmatch") if os_elem is not None else []
    fingerprinted, chosen, alt_str = evaluate_osmatches(osmatches, min_accuracy)

    if chosen is not None:
        os_name, acc_s, family, risk_category, eol_note = chosen
        return HostScanResult(
            ip=ip, hostname=hostname, subnet=subnet_label,
            fingerprinted=True, deprecated=True,
            os_name=_sanitize_text(os_name), accuracy=acc_s, family=family,
            risk_category=risk_category, eol_note=eol_note,
            alternate_guesses=alt_str, scan_status="Completed",
            notes=f"Matched deprecated-OS rule ({eol_note}).",
        )

    return HostScanResult(
        ip=ip, hostname=hostname, subnet=subnet_label,
        fingerprinted=fingerprinted, deprecated=False,
        alternate_guesses=alt_str, scan_status="Completed",
        notes=("No osmatch candidates returned by nmap -O." if not fingerprinted else
               "osmatch candidates returned, but none matched a deprecated-OS rule "
               "at/above the accuracy threshold."),
    )


_NMAP_TIME_UNIT_SECONDS = {"ms": 0.001, "s": 1.0, "m": 60.0, "h": 3600.0}


def _parse_nmap_time_spec(spec: str, default_seconds: float = 60.0) -> float:
    """Best-effort parse of an nmap time spec (e.g. '60s', '2m', '500ms',
    or a bare number of seconds) into seconds, for sizing subprocess.run's
    own timeout -- never raises; falls back to default_seconds on anything
    it doesn't recognize, since this only sizes a SAFETY NET on top of
    nmap's own --host-timeout, not the authoritative timeout value itself."""
    spec = (spec or "").strip().lower()
    m = re.match(r"^(\d+(?:\.\d+)?)(ms|s|m|h)?$", spec)
    if not m:
        return default_seconds
    value = float(m.group(1))
    unit = m.group(2) or "s"
    return value * _NMAP_TIME_UNIT_SECONDS[unit]


def _run_nmap_single_host(nmap_path: str, ip: str, xml_path: str, host_timeout: str,
                           max_rate: int) -> bool:
    """Runs nmap against exactly one host. -Pn/-n skip nmap's own host
    discovery and DNS - both already done in Phase 1 / the resolve step.
    Returns True if nmap ran (regardless of whether an OS was fingerprinted),
    False on a launch failure OR a timeout."""
    cmd = [nmap_path, "-O", "-sV", "-Pn", "-n", "--max-rate", str(max_rate),
           "--host-timeout", host_timeout, "-oX", xml_path, ip]
    # subprocess.run's own timeout is a safety net ON TOP OF --host-timeout,
    # not a replacement for it: if nmap itself hangs for a reason
    # --host-timeout doesn't cover (stuck acquiring a raw socket, a kernel/
    # driver hiccup, an orphaned child), a bare subprocess.run with no
    # timeout= blocks forever and permanently pins the worker thread -
    # Ctrl+C can't help, since Python's signal handler only runs on the main
    # thread. A generous multiplier (3x + 30s) keeps this from ever firing
    # before nmap's own --host-timeout would, under normal operation.
    hard_timeout = _parse_nmap_time_spec(host_timeout) * 3 + 30
    try:
        subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                        timeout=hard_timeout)
        return True
    except (FileNotFoundError, OSError):
        return False
    except subprocess.TimeoutExpired:
        return False


def scan_one_host(host: DiscoveredHost, args: argparse.Namespace,
                   logger: logging.Logger) -> HostScanResult:
    # Outermost guard on top of the finer-grained one below: nothing in this
    # function, including temp-file creation itself, may be allowed to
    # propagate out and take the rest of the batch down with it.
    try:
        return _scan_one_host_inner(host, args, logger)
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"scan_one_host({host.ip}) failed unexpectedly: {exc}")
        return HostScanResult(ip=host.ip, hostname=host.hostname, subnet=host.subnet,
                               fingerprinted=False, deprecated=False, scan_status="Error",
                               notes=f"scan_one_host failed unexpectedly: {exc}")


def _scan_one_host_inner(host: DiscoveredHost, args: argparse.Namespace,
                          logger: logging.Logger) -> HostScanResult:
    scan_date = datetime.now().strftime("%Y-%m-%d")
    fd, xml_path = tempfile.mkstemp(prefix="outdated_os_", suffix=".xml")
    os.close(fd)
    try:
        result: Optional[HostScanResult] = None
        for attempt in range(max(1, args.retries + 1)):
            # Broad except is deliberate, matching the sibling scripts' "one
            # bad host must not kill the scan" rule: nmap failing to even
            # write the -oX file (killed by --host-timeout mid-write, a disk
            # hiccup, a permissions problem) must not propagate past this
            # one host's result.
            try:
                ok = _run_nmap_single_host(args.nmap_path, host.ip, xml_path,
                                            args.host_timeout, args.rate)
                parsed = (parse_single_host_xml(xml_path, host.hostname, host.subnet,
                                                 args.min_accuracy) if ok else None)
            except Exception as exc:  # noqa: BLE001
                logger.debug(f"scan_one_host({host.ip}) attempt {attempt + 1} failed: {exc}")
                parsed = None
            if parsed is not None:
                result = parsed
                break
            # parsed is None here means either nmap failed to launch, or it
            # wrote no usable <host> element at all (killed by
            # --host-timeout, disk hiccup, etc). That's the "failed/empty
            # nmap invocation" case worth retrying - NOT the same as a
            # deterministic "host answered, 0 osmatch" result, which
            # parse_single_host_xml() already returns as a normal
            # (non-None) HostScanResult with fingerprinted=False, and that
            # result is never retried - see ONLY-RETRY-FAILURES rule below.
        if result is None:
            result = HostScanResult(
                ip=host.ip, hostname=host.hostname, subnet=host.subnet,
                fingerprinted=False, deprecated=False, scan_status="No Response",
                notes="nmap produced no usable host data after all retries "
                      "(launch failure, host-timeout, or unparseable XML).",
            )
        if args.keep_temp and os.path.exists(xml_path):
            # Broad except for the same reason as the retry loop above: this
            # is bookkeeping on top of an already-successful scan, and must
            # not be able to take down the rest of the batch if it fails.
            try:
                keep_dir = os.path.join(args.output_dir, f"nmap_xml_{scan_date}")
                os.makedirs(keep_dir, exist_ok=True)
                dest = os.path.join(keep_dir, f"{host.ip.replace(':', '_')}.xml")
                shutil.move(xml_path, dest)
            except Exception as exc:  # noqa: BLE001
                logger.warning(f"--keep-temp: could not save XML for {host.ip}: {exc}")
        return result
    finally:
        try:
            os.remove(xml_path)
        except OSError:
            pass


@dataclass
class Phase2Stats:
    lock: threading.Lock = field(default_factory=threading.Lock)
    total: int = 0
    completed: int = 0
    active_workers: int = 0
    fingerprinted: int = 0
    not_fingerprinted: int = 0
    risk_counts: Dict[str, int] = field(default_factory=lambda: {c: 0 for c in OS_RISK_ORDER})


def render_phase2_line(stats: Phase2Stats, start_time: float) -> str:
    with stats.lock:
        completed = stats.completed
        total = stats.total
        active = stats.active_workers
        fp = stats.fingerprinted
        nfp = stats.not_fingerprinted
        rc = dict(stats.risk_counts)
    pct = (completed / total * 100.0) if total else 0.0
    elapsed = time.time() - start_time
    if 0 < completed < total:
        eta = fmt_elapsed((elapsed / completed) * (total - completed))
    elif completed >= total and total > 0:
        eta = "00:00:00"
    else:
        eta = "n/a"
    bar = render_bar(pct)
    return (f"Phase 2 - OS FINGERPRINT {bar} | Completed: {completed}/{total} | "
            f"Workers: {active} | Fingerprinted: {fp} | Not fingerprinted: {nfp} | "
            f"Critical:{rc[OS_RISK_CRITICAL]} High:{rc[OS_RISK_HIGH]} Watch:{rc[OS_RISK_WATCH]} | "
            f"Elapsed: {fmt_elapsed(elapsed)} | ETA: {eta}")


def run_phase2(hosts: List[DiscoveredHost], args: argparse.Namespace,
               logger: logging.Logger, stop_event: threading.Event) -> List[HostScanResult]:
    stats = Phase2Stats()
    stats.total = len(hosts)
    results: List[HostScanResult] = []
    start_time = time.time()
    logger.info(f"Phase 2 - OS FINGERPRINT starting ({stats.total} hosts, {args.workers} workers)")

    def wrapped(host: DiscoveredHost) -> HostScanResult:
        with stats.lock:
            stats.active_workers += 1
        try:
            return scan_one_host(host, args, logger)
        finally:
            with stats.lock:
                stats.active_workers -= 1

    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as executor:
        futures = {executor.submit(wrapped, host): host for host in hosts}
        try:
            for future in as_completed(futures):
                result = future.result()
                with stats.lock:
                    stats.completed += 1
                    if result.fingerprinted:
                        stats.fingerprinted += 1
                    else:
                        stats.not_fingerprinted += 1
                    if result.deprecated and result.risk_category in stats.risk_counts:
                        stats.risk_counts[result.risk_category] += 1
                results.append(result)
                draw_progress_line(render_phase2_line(stats, start_time))
                if stop_event.is_set():
                    logger.warning("Interrupt received, cancelling remaining scans...")
                    for f in futures:
                        f.cancel()
                    break
        finally:
            finish_progress_line()

    logger.info(f"Phase 2 - OS FINGERPRINT complete. {stats.completed}/{stats.total} hosts "
                f"scanned, {stats.fingerprinted} fingerprinted, "
                f"finished in {fmt_elapsed(time.time() - start_time)}.")
    return results


# =============================================================================
# CSV
# =============================================================================

def write_csv_report(records: List[OutdatedOSRecord], path: str, logger: logging.Logger) -> None:
    try:
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=CSV_FIELDS)
            writer.writeheader()
            for record in records:
                row = record_to_row(record)
                for key in ("Hostname", "OS Name", "Alternate OS Guesses"):
                    row[key] = _neutralize_formula(row[key])
                writer.writerow(row)
        logger.info(f"CSV report written to {path}")
    except OSError as exc:
        logger.error(f"Failed to write CSV report to {path}: {exc}")


def _csv_field(r: dict, key: str, default: str = "") -> str:
    return r.get(key) or default


def read_rows_from_csv(csv_path: str, logger: logging.Logger) -> List[OutdatedOSRecord]:
    records: List[OutdatedOSRecord] = []
    with open(csv_path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            os_name = _sanitize_text(_csv_field(r, "OS Name"))
            # Re-derive family/risk_category from OS Name via the SAME
            # classify_deprecated()/DEPRECATED_OS_RULES table the live-scan
            # path uses, rather than trusting the CSV's stored columns
            # verbatim. Without this, editing DEPRECATED_OS_RULES later
            # (e.g. bumping a keyword's severity after its EOL date passes)
            # and reloading an old CSV via --from-csv would keep showing the
            # stale classification forever, silently diverging from what a
            # fresh live re-scan of the same hosts would now report.
            rule = classify_deprecated(os_name)
            if rule is not None:
                os_family, raw_risk_category, _eol_note = rule
            else:
                # OS Name doesn't match any current rule (hand-edited CSV,
                # blank field, or a keyword removed from the table since this
                # CSV was written) - this report's CSV only ever contains
                # deprecated-host rows (no "clean"/"not fingerprinted" bucket
                # exists in OS_RISK_ORDER to fall back to), so fall back to
                # the CSV's own stored columns, validating Risk Category
                # against OS_RISK_ORDER the same way as before.
                os_family = _csv_field(r, "OS Family", "Other")
                raw_risk_category = _csv_field(r, "Risk Category", OS_RISK_HIGH)
                normalized = normalize_risk_category(raw_risk_category)
                if normalized != raw_risk_category:
                    logger.warning(
                        f"Unrecognized Risk Category '{raw_risk_category}' for "
                        f"{_csv_field(r, 'IP Address')} in {csv_path}; defaulting to '{normalized}'.")
                raw_risk_category = normalized
            records.append(OutdatedOSRecord(
                scan_date=_csv_field(r, "Scan Date"), ip=_csv_field(r, "IP Address"),
                hostname=_sanitize_text(_csv_field(r, "Hostname")),
                subnet=_csv_field(r, "Subnet Range", "UNKNOWN"),
                os_name=os_name,
                os_family=os_family,
                accuracy=_csv_field(r, "Accuracy", "0"),
                risk_category=raw_risk_category,
                alternate_guesses=_sanitize_text(_csv_field(r, "Alternate OS Guesses"), max_len=1500),
                notes=_sanitize_text(_csv_field(r, "Notes")),
                meaning=OS_RISK_MEANING.get(raw_risk_category, _sanitize_text(_csv_field(r, "Meaning"))),
                scan_status=_csv_field(r, "Scan Status", "Completed"),
            ))
    return records


# =============================================================================
# XLSX
# =============================================================================

def compute_category_stats(records: List[OutdatedOSRecord]) -> Dict[str, int]:
    counts = {c: 0 for c in OS_RISK_ORDER}
    for rec in records:
        counts[normalize_risk_category(rec.risk_category)] += 1
    return counts


def breakdown_by_field(records: List[OutdatedOSRecord], key_func
                        ) -> List[Tuple[str, int, float, int]]:
    """Groups records by key_func(record); returns (label, count, avg_acc,
    min_acc) rows sorted by count descending, then label. Restores the old
    outdatedOS.py script's per-OS-family/per-OS-name accuracy breakdown
    (dropped during the house-convention rewrite with no replacement) -
    this is exactly the 'accurate metrics' granularity a reviewer uses to
    triage which OS/accuracy clusters to prioritize first."""
    buckets: Dict[str, List[int]] = {}
    for rec in records:
        label = key_func(rec)
        try:
            acc = int(rec.accuracy)
        except (ValueError, TypeError):
            acc = 0
        buckets.setdefault(label, []).append(acc)
    rows = [(label, len(accs), round(sum(accs) / len(accs), 1), min(accs))
            for label, accs in buckets.items()]
    rows.sort(key=lambda row: (-row[1], row[0]))
    return rows


def build_workbook(records: List[OutdatedOSRecord], networks: List[ipaddress.IPv4Network],
                    scan_stats: Optional[Dict[str, int]]):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    def fill(hex_color: str):
        return PatternFill("solid", fgColor=hex_color)

    def thin_border():
        s = Side(style="thin", color="D5D8DC")
        return Border(left=s, right=s, top=s, bottom=s)

    wb = Workbook()

    # --- Overview sheet ---
    ov = wb.active
    ov.title = "Overview"
    ov.sheet_view.showGridLines = False
    ov.column_dimensions["A"].width = 34
    ov.column_dimensions["B"].width = 70
    ov.column_dimensions["C"].width = 14

    ov.append(["Outdated / Deprecated OS Scan"])
    ov["A1"].font = Font(bold=True, size=14)
    ov.append([f"Scan date: {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    ov.append([f"Configured subnets in scope: {len(networks)}"])
    ov.append([f"Masscan probe ports (Phase 1): {MASSCAN_PORTS}"])
    ov.append([f"Minimum osmatch accuracy threshold: {MIN_ACCURACY}%"])
    ov.append([])

    ov.append(["How to read this report"])
    ov[f"A{ov.max_row}"].font = Font(bold=True, size=12)
    ov.append(["This report only lists hosts where nmap's OS fingerprint matched a known "
               "deprecated/end-of-life OS at or above the accuracy threshold - there is no "
               "'clean' row for every scanned host (see Scan Summary below for the full host "
               "counts, including hosts that were scanned but NOT fingerprinted at all). Every "
               "listed host falls into exactly one of three severity categories - see the table "
               "below."])
    ov.cell(row=ov.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ov.merge_cells(start_row=ov.max_row, start_column=1, end_row=ov.max_row, end_column=3)
    ov.row_dimensions[ov.max_row].height = 70
    ov.append([])

    ov.append(["Category", "What it means"])
    for cell in ov[ov.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill("2C3E50")
    for category in OS_RISK_ORDER:
        ov.append([f"{OS_RISK_EMOJI[category]} {category}", OS_RISK_DESCRIPTION[category]])
        row = ov.max_row
        ov.cell(row=row, column=1).fill = fill(OS_RISK_FILL_HEX[category])
        ov.cell(row=row, column=1).font = Font(color=OS_RISK_FONT_HEX[category], bold=True)
        ov.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ov.row_dimensions[row].height = 60
    ov.append([])

    ov.append(["Methodology notes"])
    ov[f"A{ov.max_row}"].font = Font(bold=True, size=12)
    for note in [
        "Phase 1 (masscan) finds every host with one of the configured probe ports open. Phase 2 "
        "runs one 'nmap -O -sV -Pn -n' process per discovered host and parses nmap's XML "
        "<os><osmatch> candidates.",
        "FIX: this script checks EVERY osmatch candidate nmap returns (already sorted by "
        "accuracy, descending), not just nmap's single best guess. A host whose top guess isn't "
        "a deprecated OS but whose 2nd- or 3rd-best guess (still above the accuracy threshold) IS "
        "one would previously have been silently missed - undercounting deprecated hosts. Every "
        "osmatch examined is kept in the 'Alternate OS Guesses' column as evidence.",
        "nmap -O/-sV only sends standard TCP/IP-stack and service probe packets to infer OS/"
        "service identity from response behavior. No authentication is attempted, no "
        "exploitation, no data modification.",
    ]:
        ov.append([note])
        ov.cell(row=ov.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ov.merge_cells(start_row=ov.max_row, start_column=1, end_row=ov.max_row, end_column=3)
        ov.row_dimensions[ov.max_row].height = 55
    ov.append([])

    # --- Scan Summary ---
    total = len(records)
    counts = compute_category_stats(records)
    ov.append(["Scan Summary"])
    ov[f"A{ov.max_row}"].font = Font(bold=True, size=12)

    def stat_or_na(key: str) -> object:
        if scan_stats is None or scan_stats.get(key) is None:
            return "N/A (not captured when rebuilding from --from-csv)"
        return scan_stats[key]

    ov.append(["Live hosts discovered in Phase 1:", stat_or_na("live_hosts")])
    ov.append(["Hosts fingerprinted (returned >=1 osmatch):", stat_or_na("fingerprinted")])
    ov.append(["Hosts NOT fingerprinted (inconclusive):", stat_or_na("not_fingerprinted")])
    ov.append(["Total hosts scanned in Phase 2:", stat_or_na("total_hosts_scanned")])
    ov.append([])
    ov.append(["Total deprecated-OS hosts found:", total])
    ov.append(["Category", "Count", "% of Total"])
    for cell in ov[ov.max_row]:
        cell.font = Font(bold=True)
    for category in OS_RISK_ORDER:
        pct = (counts[category] / total * 100.0) if total else 0.0
        ov.append([f"{OS_RISK_EMOJI[category]} {category}", counts[category], round(pct, 1)])
        ov.cell(row=ov.max_row, column=1).fill = fill(OS_RISK_FILL_HEX[category])
        ov.cell(row=ov.max_row, column=3).number_format = "0.0"
    ov.append([])

    # Subnet breakdown (of the deprecated-host rows only - same grain as the
    # Scan Results sheet)
    subnet_counts: Dict[str, int] = {}
    for r in records:
        subnet_counts[r.subnet] = subnet_counts.get(r.subnet, 0) + 1
    ov.append(["Subnet Breakdown (deprecated-OS hosts)"])
    ov[f"A{ov.max_row}"].font = Font(bold=True, size=12)
    for subnet, count in sorted(subnet_counts.items()):
        ov.append([subnet, count])
    ov.append([])

    # OS family / name accuracy breakdown - the per-cluster granularity
    # that makes the "how many, how confident" metrics actually actionable
    # (which family/version to prioritize first), not just a flat host list.
    ov.append(["OS Family Breakdown (deprecated-OS hosts)"])
    ov[f"A{ov.max_row}"].font = Font(bold=True, size=12)
    ov.append(["OS Family", "Hosts", "Avg Accuracy %", "Min Accuracy %"])
    for cell in ov[ov.max_row]:
        cell.font = Font(bold=True)
    for label, count, avg_acc, min_acc in breakdown_by_field(records, lambda r: r.os_family):
        ov.append([label, count, avg_acc, min_acc])
    ov.append([])

    ov.append(["OS Name Breakdown (deprecated-OS hosts)"])
    ov[f"A{ov.max_row}"].font = Font(bold=True, size=12)
    ov.append(["OS Name", "Hosts", "Avg Accuracy %", "Min Accuracy %"])
    for cell in ov[ov.max_row]:
        cell.font = Font(bold=True)
    for label, count, avg_acc, min_acc in breakdown_by_field(records, lambda r: r.os_name):
        ov.append([label, count, avg_acc, min_acc])

    # --- Scan Results sheet ---
    wr = wb.create_sheet("Scan Results")
    wr.sheet_view.showGridLines = False
    wr.freeze_panes = "A2"

    widths = {"Scan Date": 12, "IP Address": 16, "Hostname": 26, "Subnet Range": 18,
              "OS Name": 34, "OS Family": 18, "Accuracy": 10, "Risk Category": 36,
              "Alternate OS Guesses": 50, "Notes": 40, "Meaning": 50, "Scan Status": 14}
    wrap_cols = {"OS Name", "Alternate OS Guesses", "Notes", "Meaning"}

    for col_idx, col_name in enumerate(CSV_FIELDS, start=1):
        c = wr.cell(row=1, column=col_idx, value=col_name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill("2C3E50")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border()
        wr.column_dimensions[get_column_letter(col_idx)].width = widths.get(col_name, 16)
    wr.row_dimensions[1].height = 22

    risk_col = CSV_FIELDS.index("Risk Category") + 1
    for row_idx, record in enumerate(records, start=2):
        row = record_to_row(record)
        for col_idx, key in enumerate(CSV_FIELDS, start=1):
            val = row[key]
            if key in ("Hostname", "OS Name", "Alternate OS Guesses"):
                val = _neutralize_formula(val)
            c = wr.cell(row=row_idx, column=col_idx, value=val)
            c.border = thin_border()
            c.alignment = Alignment(vertical="top", wrap_text=(key in wrap_cols), horizontal="left")
        cat = normalize_risk_category(record.risk_category)
        cell = wr.cell(row=row_idx, column=risk_col)
        cell.value = f"{OS_RISK_EMOJI[cat]} {cat}"
        for col_idx in range(1, len(CSV_FIELDS) + 1):
            wr.cell(row=row_idx, column=col_idx).fill = fill(OS_RISK_FILL_HEX[cat])
        wr.row_dimensions[row_idx].height = 40

    if records:
        wr.auto_filter.ref = f"A1:{get_column_letter(len(CSV_FIELDS))}{len(records) + 1}"

    return wb


def write_xlsx_report(records: List[OutdatedOSRecord], networks: List[ipaddress.IPv4Network],
                       scan_stats: Optional[Dict[str, int]], path: str,
                       logger: logging.Logger) -> None:
    try:
        wb = build_workbook(records, networks, scan_stats)
    except ImportError:
        logger.warning("openpyxl not installed - skipping .xlsx generation. "
                        "Install with: pip install openpyxl")
        return
    except Exception as exc:  # noqa: BLE001
        logger.error(f"Failed to build XLSX workbook: {exc}")
        return
    try:
        wb.save(path)
        logger.info(f"XLSX report written to {path} ({len(records)} host(s))")
    except Exception as exc:  # noqa: BLE001
        logger.error(f"Failed to write XLSX report to {path}: {exc}")


# =============================================================================
# MAIN
# =============================================================================

def parse_args(argv: List[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Authorized internal deprecated/EOL-OS fingerprint assessment.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    parser.add_argument("--workers", type=int, default=DEFAULT_WORKERS)
    parser.add_argument("--rate", type=int, default=DEFAULT_RATE,
                         help="masscan packets/sec (Phase 1) and nmap --max-rate ceiling (Phase 2)")
    parser.add_argument("--host-timeout", type=str, default=DEFAULT_HOST_TIMEOUT,
                         help="nmap --host-timeout per host, so one hung host can't stall a worker")
    parser.add_argument("--retries", type=int, default=DEFAULT_RETRIES)
    parser.add_argument("--output-dir", type=str, default=SCRIPT_DIR)
    parser.add_argument("--masscan-path", type=str, default="masscan")
    parser.add_argument("--nmap-path", type=str, default="nmap")
    parser.add_argument("--interface", type=str, default=None)
    parser.add_argument("--skip-masscan", action="store_true")
    parser.add_argument("--masscan-output-file", type=str, default=None)
    parser.add_argument("--no-xlsx", action="store_true")
    parser.add_argument("--csv-out", type=str, default=None)
    parser.add_argument("--xlsx-out", type=str, default=None)
    parser.add_argument("--from-csv", metavar="FILE",
                         help="Skip scanning entirely; rebuild the .xlsx from an existing CSV")
    parser.add_argument("--min-accuracy", type=int, default=MIN_ACCURACY,
                         help="Ignore osmatch guesses below this confidence (%%)")
    parser.add_argument("--keep-temp", action="store_true",
                         help="keep each host's raw nmap XML under <output-dir>/nmap_xml_<date>/")
    return parser.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    args = parse_args(argv if argv is not None else sys.argv[1:])

    scan_date = datetime.now().strftime("%Y-%m-%d")
    try:
        os.makedirs(args.output_dir, exist_ok=True)
    except OSError as exc:
        print(f"ERROR: Could not create output directory '{args.output_dir}': {exc}", file=sys.stderr)
        return 1

    log_path = os.path.join(args.output_dir, f"outdated_os_scan_{scan_date}.log")
    csv_path = args.csv_out or os.path.join(args.output_dir, f"outdated_os_scan_{scan_date}.csv")
    xlsx_path = args.xlsx_out or os.path.join(args.output_dir, f"outdated_os_scan_{scan_date}.xlsx")

    try:
        logger = setup_logging(log_path)
    except OSError as exc:
        print(f"ERROR: Could not open log file '{log_path}': {exc}", file=sys.stderr)
        return 1

    networks = validate_subnets(SUBNETS, logger)

    if args.from_csv:
        try:
            records = read_rows_from_csv(args.from_csv, logger)
        except OSError as exc:
            logger.error(f"Could not read --from-csv file '{args.from_csv}': {exc}")
            return 1
        except (csv.Error, UnicodeDecodeError) as exc:
            logger.error(f"--from-csv file '{args.from_csv}' is not a valid UTF-8 CSV: {exc}")
            return 1
        if not args.no_xlsx:
            write_xlsx_report(records, networks, None, xlsx_path, logger)
        return 0

    stop_event = threading.Event()

    def handle_sigint(signum, frame):  # noqa: ANN001
        if stop_event.is_set():
            logger.warning("Second interrupt received, forcing exit.")
            sys.exit(130)
        logger.warning("Ctrl+C received - finishing current work and writing "
                        "partial reports. Press Ctrl+C again to force exit.")
        stop_event.set()

    signal.signal(signal.SIGINT, handle_sigint)

    logger.info("=" * 70)
    logger.info("AUTHORIZED OUTDATED/DEPRECATED OS SCAN")
    logger.info("=" * 70)
    logger.info("Configured subnets in scope:")
    for s in SUBNETS:
        logger.info(f"  - {s}")
    logger.info(f"Masscan probe ports: {MASSCAN_PORTS}")
    logger.info(f"Workers: {args.workers} | Rate: {args.rate} | "
                f"Host timeout: {args.host_timeout} | Retries: {args.retries} | "
                f"Min accuracy: {args.min_accuracy}%")

    if not networks:
        logger.error("No valid subnets configured. Exiting.")
        return 1

    masscan_path = check_external_tool(args.masscan_path) or args.masscan_path
    if not args.skip_masscan and check_external_tool(args.masscan_path) is None:
        logger.error(f"Required tool '{args.masscan_path}' was not found on PATH.")
        return 1

    resolved_nmap = check_external_tool(args.nmap_path)
    if resolved_nmap is None:
        logger.error(f"Required tool '{args.nmap_path}' was not found on PATH. "
                      f"There is no fallback for Phase 2 - nmap -O/-sV is the only "
                      f"OS-fingerprint engine this script has.")
        return 1
    args.nmap_path = resolved_nmap

    raw_records: List[Tuple[str, int, str]] = []
    if args.skip_masscan:
        if not args.masscan_output_file or not os.path.exists(args.masscan_output_file):
            logger.error("--skip-masscan requires a valid --masscan-output-file.")
            return 1
        raw_records, _ = parse_masscan_list_output(args.masscan_output_file, 0)
    else:
        masscan_out = args.masscan_output_file or os.path.join(
            args.output_dir, f".masscan_output_{scan_date}.txt")
        try:
            raw_records = run_masscan_phase1(masscan_path, SUBNETS, args.rate, masscan_out,
                                              args.interface, parse_port_spec(MASSCAN_PORTS),
                                              logger, stop_event)
        except Exception as exc:  # noqa: BLE001
            logger.error(f"Phase 1 discovery failed unexpectedly: {exc}")
            raw_records = []

    dedup: set = set()
    for ip, _port, _proto in raw_records:
        try:
            ipaddress.ip_address(ip)
        except ValueError:
            logger.warning(f"Skipping malformed IP address from scan output: {ip}")
            continue
        dedup.add(ip)

    live_hosts_count = len(dedup)
    logger.info(f"Discovered {live_hosts_count} unique live host(s) after deduplication.")

    hosts: List[DiscoveredHost] = []
    for ip in sorted(dedup, key=lambda x: ipaddress.ip_address(x)):
        if stop_event.is_set():
            break
        hostname = resolve_hostname(ip, timeout=2.0)
        subnet = subnet_for_ip(ip, networks)
        hosts.append(DiscoveredHost(ip=ip, hostname=hostname, subnet=subnet))

    results: List[HostScanResult] = []
    if hosts:
        # Deliberately not gated on "and not stop_event.is_set()" - same
        # reasoning as the sibling scripts: a Ctrl+C during the hostname-
        # resolution loop above can set stop_event while hosts is already
        # non-empty, and run_phase2() itself already honors stop_event
        # correctly (cancels remaining futures, returns whatever completed).
        results = run_phase2(hosts, args, logger, stop_event)
    else:
        logger.info("No live hosts discovered; skipping Phase 2 scan.")

    fingerprinted_count = sum(1 for r in results if r.fingerprinted)
    not_fingerprinted_count = len(results) - fingerprinted_count

    deprecated_results = [r for r in results if r.deprecated]
    records = [
        OutdatedOSRecord(
            scan_date=scan_date, ip=r.ip, hostname=r.hostname, subnet=r.subnet,
            os_name=r.os_name, os_family=r.family, accuracy=r.accuracy,
            risk_category=r.risk_category, alternate_guesses=r.alternate_guesses,
            notes=r.notes, meaning=OS_RISK_MEANING.get(r.risk_category, ""),
            scan_status=r.scan_status,
        )
        for r in deprecated_results
    ]
    records.sort(key=lambda r: ipaddress.ip_address(r.ip))

    scan_stats = {
        "live_hosts": live_hosts_count,
        "fingerprinted": fingerprinted_count,
        "not_fingerprinted": not_fingerprinted_count,
        "total_hosts_scanned": len(results),
    }

    write_csv_report(records, csv_path, logger)
    if not args.no_xlsx:
        write_xlsx_report(records, networks, scan_stats, xlsx_path, logger)

    counts = compute_category_stats(records)
    logger.info("=" * 70)
    logger.info("SCAN SUMMARY")
    logger.info(f"  Live hosts: {live_hosts_count} | Fingerprinted: {fingerprinted_count} | "
                f"Not fingerprinted: {not_fingerprinted_count}")
    logger.info(f"  Deprecated-OS hosts found: {len(records)} | "
                f"Critical: {counts[OS_RISK_CRITICAL]} | "
                f"High: {counts[OS_RISK_HIGH]} | "
                f"Watch: {counts[OS_RISK_WATCH]}")
    logger.info("=" * 70)
    logger.info(f"Reports written to: {os.path.abspath(args.output_dir)}")

    if stop_event.is_set():
        logger.warning("Scan was interrupted by user; reports reflect partial results.")
        return 130
    return 0


if __name__ == "__main__":
    sys.exit(main())
